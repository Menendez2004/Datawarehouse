"""Genera el documento de la Fase 1 (propuesta y diseño del Datawarehouse).

Uso:  python docs/build_doc.py ["ruta/SAC Retailer Dataset.xlsx"]
Salida: diagrams/*.png y output/DMD941_Proyecto_Fase_1.docx
"""
import collections
import datetime as dt
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import openpyxl
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT = Path(__file__).resolve().parent.parent
REPO = PROJECT.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO / "SAC Retailer Dataset.xlsx"
DIAGRAMS = PROJECT / "diagrams"
OUTPUT = PROJECT / "output"
LOGO = REPO / "img" / "udb-virtual.png"

NAVY = RGBColor(0x1F, 0x38, 0x64)
ACCENT_HEX = "2C5F8A"
LIGHT_HEX = "F2F5F8"

INTEGRANTES = [
    ("Kevin Edenilson Menéndez Zepeda", "MZ232727"),
]
DOCENTE = "Ing. David Enrique Moya"
FECHA_ENTREGA = "23 de octubre de 2026"

COSTOS = ["Shipping Costs", "Storage Costs", "Personnel Cost", "Selling Costs", "Purchasing Costs"]


# ================= Perfilado del Excel =================
def perfilar(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header, data = rows[0], [dict(zip(rows[0], r)) for r in rows[1:]]

    def anual(year):
        y = [r for r in data if r["Order Date"].year == year]
        ventas = sum(r["Sales"] for r in y)
        utilidad = sum(r["Profit"] for r in y)
        pedidos = len({(r["Order ID"], r["Customer ID"], r["Order Date"]) for r in y})
        clientes = len({r["Customer ID"] for r in y})
        return dict(ventas=ventas, utilidad=utilidad, pedidos=pedidos, clientes=clientes,
                    unidades=sum(r["Quantity"] for r in y))

    ventas_tot = sum(r["Sales"] for r in data)
    pedidos_por_id = collections.defaultdict(set)
    for r in data:
        pedidos_por_id[r["Order ID"]].add((r["Customer ID"], r["Order Date"]))
    regiones = collections.defaultdict(set)
    for r in data:
        regiones[r["Region"]].add(r["Market"])

    return dict(
        columnas=list(header),
        filas=len(data),
        desde=min(r["Order Date"] for r in data),
        hasta=max(r["Order Date"] for r in data),
        clientes=len({r["Customer ID"] for r in data}),
        productos=len({r["Product ID"] for r in data}),
        ciudades=len({(r["City"], r["State"], r["Country"]) for r in data}),
        paises=len({r["Country"] for r in data}),
        mercados=sorted({r["Market"] for r in data}),
        categorias=sorted({r["Category"] for r in data}),
        subcategorias=len({r["Sub-Category"] for r in data}),
        modos=sorted({r["Ship Mode"] for r in data}),
        envio_antes=sum(r["Ship Date"] < r["Order Date"] for r in data),
        perdidas=sum(r["Profit"] < 0 for r in data),
        order_id_reusado=sum(len(v) > 1 for v in pedidos_por_id.values()),
        regiones_ambiguas=sorted(k for k, v in regiones.items() if len(v) > 1),
        descuentos_flotantes=sum(1 for r in data if round(r["Discount"], 2) != r["Discount"]),
        utilidad_cuadra=max(abs(r["Sales"] - sum(r[c] for c in COSTOS) - r["Profit"]) for r in data) < 0.01,
        pct_costos={c: sum(r[c] for r in data) / ventas_tot * 100 for c in COSTOS},
        y2019=anual(2019),
        y2020=anual(2020),
    )


# ================= Diagramas =================
TABLAS = {
    "dim_fecha": ["PK fecha_key (AAAAMMDD)", "fecha", "dia", "mes", "nombre_mes", "trimestre",
                  "anio", "dia_semana", "nombre_dia", "es_fin_semana"],
    "dim_cliente": ["PK cliente_key", "customer_id", "nombre_cliente", "segmento", "cluster_kmeans"],
    "dim_producto": ["PK producto_key", "product_id", "nombre_producto", "subcategoria", "categoria"],
    "dim_geografia": ["PK geografia_key", "ciudad", "estado", "pais", "iso3", "region", "mercado",
                      "latitud", "longitud"],
    "dim_modo_envio": ["PK modo_envio_key", "modo_envio"],
    "fact_ventas": ["PK venta_key", "primary_id", "order_id (degenerada)", "FK fecha_pedido_key",
                    "FK fecha_envio_key", "FK cliente_key", "FK producto_key", "FK geografia_key",
                    "FK modo_envio_key", "cantidad", "descuento", "ventas", "costo_envio",
                    "costo_almacenamiento", "costo_personal", "costo_venta", "costo_compra",
                    "costo_total", "utilidad", "dias_envio", "fecha_inconsistente"],
}


def dibujar_tabla(ax, nombre, x, y, ancho, color):
    filas = TABLAS[nombre]
    alto_fila, alto_tit = 0.42, 0.55
    alto = alto_tit + alto_fila * len(filas)
    ax.add_patch(plt.Rectangle((x, y - alto), ancho, alto, fc="white", ec=color, lw=1.6, zorder=2))
    ax.add_patch(plt.Rectangle((x, y - alto_tit), ancho, alto_tit, fc=color, ec=color, zorder=3))
    ax.text(x + ancho / 2, y - alto_tit / 2, nombre, ha="center", va="center", color="white",
            fontsize=10, fontweight="bold", zorder=4)
    for i, campo in enumerate(filas):
        peso = "bold" if campo.startswith(("PK", "FK")) else "normal"
        ax.text(x + 0.12, y - alto_tit - alto_fila * (i + 0.5), campo, va="center", fontsize=8,
                fontweight=peso, zorder=4)
    return (x, y - alto, ancho, alto)


def diagrama_estrella(path):
    fig, ax = plt.subplots(figsize=(13, 9.5))
    ax.set_xlim(0, 26)
    ax.set_ylim(0, 19)
    ax.axis("off")
    hecho, dim = "#1F3864", "#2C5F8A"
    cajas = {
        "fact_ventas": dibujar_tabla(ax, "fact_ventas", 10.5, 14.3, 5.0, hecho),
        "dim_fecha": dibujar_tabla(ax, "dim_fecha", 0.5, 18.5, 4.6, dim),
        "dim_cliente": dibujar_tabla(ax, "dim_cliente", 20.8, 18.5, 4.6, dim),
        "dim_producto": dibujar_tabla(ax, "dim_producto", 20.8, 8.5, 4.6, dim),
        "dim_geografia": dibujar_tabla(ax, "dim_geografia", 0.5, 7.5, 4.6, dim),
        "dim_modo_envio": dibujar_tabla(ax, "dim_modo_envio", 10.9, 3.0, 4.2, dim),
    }
    fx, fy, fw, fh = cajas["fact_ventas"]
    relaciones = [
        ("dim_fecha", (fx, fy + fh - 2.2), "fecha_pedido_key / fecha_envio_key"),
        ("dim_fecha", (fx, fy + fh - 2.6), None),
        ("dim_cliente", (fx + fw, fy + fh - 3.0), "cliente_key"),
        ("dim_producto", (fx + fw, fy + fh - 3.4), "producto_key"),
        ("dim_geografia", (fx, fy + fh - 3.8), "geografia_key"),
        ("dim_modo_envio", (fx + fw / 2, fy), "modo_envio_key"),
    ]
    for dim_nombre, (px, py), etiqueta in relaciones:
        x, y, w, h = cajas[dim_nombre]
        if dim_nombre == "dim_modo_envio":
            dx, dy = x + w / 2, y + h
        else:
            dx = x + w if x < fx else x
            dy = y + h - 0.75
        ax.annotate("", xy=(dx, dy), xytext=(px, py),
                    arrowprops=dict(arrowstyle="-", color="#555555", lw=1.2), zorder=1)
        if etiqueta is None:
            continue
        ax.text(px + (dx - px) * 0.5, py + (dy - py) * 0.5 + 0.18, f"1 : N  ({etiqueta})",
                fontsize=7.5, ha="center", color="#333333", style="italic",
                bbox=dict(fc="white", ec="none", pad=0.5), zorder=5)
    ax.set_title("Global Electronic Retail — Esquema estrella del Datawarehouse",
                 fontsize=13, fontweight="bold", color=hecho)
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


CRONOGRAMA = [
    # (actividad, inicio, fin, responsable)
    ("Configuración del entorno (PostgreSQL, Power BI, repositorio)", "2026-10-27", "2026-10-29", "Kevin Menéndez"),
    ("Creación del esquema del DW (script DDL)", "2026-10-30", "2026-11-02", "Kevin Menéndez"),
    ("ETL: extracción y perfilado del Excel", "2026-11-03", "2026-11-05", "Kevin Menéndez"),
    ("ETL: transformación y limpieza", "2026-11-06", "2026-11-11", "Kevin Menéndez"),
    ("ETL: carga y validación (conteos y sumas)", "2026-11-12", "2026-11-14", "Kevin Menéndez"),
    ("Modelo de Power BI y medidas DAX de los KPIs", "2026-11-15", "2026-11-20", "Kevin Menéndez"),
    ("Segmentación de clientes con K-means", "2026-11-18", "2026-11-23", "Kevin Menéndez"),
    ("Dashboard interactivo de finanzas y ventas", "2026-11-21", "2026-11-27", "Kevin Menéndez"),
    ("Pruebas, validación de KPIs y correcciones", "2026-11-28", "2026-12-01", "Kevin Menéndez"),
    ("Documentación final y presentación", "2026-12-02", "2026-12-06", "Kevin Menéndez"),
]


def diagrama_gantt(path):
    fig, ax = plt.subplots(figsize=(12, 5.2))
    for i, (act, ini, fin, _) in enumerate(CRONOGRAMA):
        a = dt.date.fromisoformat(ini)
        b = dt.date.fromisoformat(fin) + dt.timedelta(days=1)
        ax.barh(i, (b - a).days, left=mdates.date2num(a), color="#2C5F8A", height=0.55)
    ax.set_yticks(range(len(CRONOGRAMA)))
    ax.set_yticklabels([c[0] for c in CRONOGRAMA], fontsize=9)
    ax.invert_yaxis()
    ax.xaxis_date()
    ax.xaxis.set_major_locator(mdates.WeekdayLocator(byweekday=mdates.MO))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d/%m"))
    ax.grid(axis="x", color="#DDDDDD")
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_title("Cronograma de la fase de implementación", fontsize=12, fontweight="bold", color="#1F3864")
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


# ================= Helpers de Word =================
def sombrear(celda, color):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def tabla(doc, encabezados, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(encabezados):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(c, ACCENT_HEX)
    for ri, fila in enumerate(filas):
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            celdas[i].paragraphs[0].add_run(str(v)).font.size = Pt(9.5)
            if ri % 2:
                sombrear(celdas[i], LIGHT_HEX)
    if anchos:
        for fila in t.rows:
            for i, w in enumerate(anchos):
                fila.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def parrafo(doc, texto, negrita_inicial=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if negrita_inicial:
        p.add_run(negrita_inicial).bold = True
    p.add_run(texto)
    return p


def viñeta(doc, texto, negrita_inicial=None):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if negrita_inicial:
        p.add_run(negrita_inicial).bold = True
    p.add_run(texto)
    return p


def numerado(doc, texto, negrita_inicial=None):
    p = doc.add_paragraph(style="List Number")
    if negrita_inicial:
        p.add_run(negrita_inicial).bold = True
    p.add_run(texto)
    return p


def dinero(v):
    return f"${v:,.2f}"


# ================= Documento =================
def construir(perfil):
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    for nivel, tam in ((1, 16), (2, 13)):
        h = doc.styles[f"Heading {nivel}"]
        h.font.name = "Calibri"
        h.font.size = Pt(tam)
        h.font.color.rgb = NAVY
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Inches(1))

    # ---------- Portada ----------
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(3.2))

    def centrado(texto, tam, negrita=False, color=None, espacio=6, cursiva=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(espacio)
        r = p.add_run(texto)
        r.font.size, r.bold, r.italic = Pt(tam), negrita, cursiva
        if color:
            r.font.color.rgb = color
        return p

    centrado("UNIVERSIDAD DON BOSCO", 20, True, NAVY)
    centrado("Facultad de Ingeniería", 13)
    centrado("Datawarehouse y Minería de Datos (DMD941)", 13, espacio=24)
    centrado("Proyecto de Cátedra — Fase 1", 18, True, NAVY)
    centrado("Propuesta del proyecto y diseño del Datawarehouse", 13, cursiva=True)
    centrado("Caso: Global Electronic Retail (SAC Retailer Dataset)", 12, espacio=28)
    centrado("Integrantes" if len(INTEGRANTES) > 1 else "Alumno", 12, True)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(("Nombre completo", "Carné")):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(c, ACCENT_HEX)
    for nombre, carne in INTEGRANTES:
        c = t.add_row().cells
        c[0].text, c[1].text = nombre, carne
    for fila in t.rows:
        fila.cells[0].width, fila.cells[1].width = Inches(3.8), Inches(1.6)
    doc.add_paragraph()
    centrado(f"Docente: {DOCENTE}", 12)
    centrado("Experiencia de Aprendizaje 2 — Ponderación: 30 %", 11)
    centrado(f"Fecha de entrega: {FECHA_ENTREGA}", 11)
    centrado("San Salvador, El Salvador", 11)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    y19, y20 = perfil["y2019"], perfil["y2020"]
    margen19 = y19["utilidad"] / y19["ventas"] * 100
    margen20 = y20["utilidad"] / y20["ventas"] * 100
    crec = (y20["ventas"] / y19["ventas"] - 1) * 100

    # ---------- 1. Información general y problema ----------
    doc.add_heading("1. Información general y definición del problema", 1)
    doc.add_heading("1.1 Contexto de la empresa", 2)
    parrafo(doc, "Global Electronic Retail es un minorista internacional de productos electrónicos que "
            f"comercializa artículos de {len(perfil['categorias'])} categorías ({', '.join(perfil['categorias'])}) "
            f"y {perfil['subcategorias']} subcategorías en {perfil['paises']} países, agrupados en "
            f"{len(perfil['mercados'])} mercados ({', '.join(perfil['mercados'])}). Atiende a clientes de los "
            "segmentos Consumer (consumidor final) y Corporate (empresas), y ofrece cuatro modalidades de "
            f"envío ({', '.join(perfil['modos'])}).")
    parrafo(doc, f"Toda su información transaccional está concentrada en un único archivo de Excel "
            f"(«SAC Retailer Dataset.xlsx») con {perfil['filas']:,} líneas de pedido registradas entre el "
            f"{perfil['desde']:%d/%m/%Y} y el {perfil['hasta']:%d/%m/%Y}. Cada línea combina datos del "
            "cliente, del producto, de la ubicación geográfica, del envío y cinco rubros de costo, además de "
            "la venta y la utilidad.")

    doc.add_heading("1.2 Necesidad de las áreas de finanzas y ventas", 2)
    viñeta(doc, " necesita saber dónde se gana y dónde se pierde dinero: qué productos, mercados y "
           "niveles de descuento generan utilidad negativa, y cuál de los cinco rubros de costo (envío, "
           "almacenamiento, personal, venta y compra) erosiona más el margen.", "Finanzas:")
    viñeta(doc, " necesita medir el crecimiento de las ventas, el valor de cada pedido y de cada "
           "cliente, y comparar el desempeño por mercado, categoría, segmento y periodo para orientar "
           "campañas y metas comerciales.", "Ventas:")

    doc.add_heading("1.3 Planteamiento del problema", 2)
    parrafo(doc, "Actualmente el análisis se hace directamente sobre la hoja de cálculo. Esto presenta "
            "los siguientes problemas:")
    viñeta(doc, f" con {perfil['filas']:,} filas y 29 columnas, las tablas dinámicas se vuelven lentas y "
           "cada analista construye sus propios cálculos, lo que produce cifras distintas para un mismo "
           "indicador.", "Falta de una única versión de la verdad:")
    viñeta(doc, f" {perfil['envio_antes']} líneas tienen fecha de envío anterior a la fecha de pedido, "
           f"{perfil['order_id_reusado']:,} números de pedido se repiten con clientes o fechas distintas y "
           f"las regiones {', '.join(perfil['regiones_ambiguas'])} existen en más de un mercado.",
           "Problemas de calidad de datos:")
    viñeta(doc, " no existe un tablero que muestre los KPIs de forma periódica; las decisiones se toman "
           "con reportes manuales y desactualizados.", "Ausencia de indicadores:")
    viñeta(doc, f" {perfil['perdidas']:,} líneas de pedido ({perfil['perdidas'] / perfil['filas'] * 100:.1f} %) "
           "generan pérdida, pero el formato plano no permite identificar fácilmente sus causas.",
           "Pérdidas ocultas:")

    doc.add_heading("1.4 Solución propuesta y aporte del Datawarehouse", 2)
    parrafo(doc, "Se propone construir un Datawarehouse en PostgreSQL con un modelo dimensional de "
            "esquema estrella, alimentado mediante un proceso ETL que extraiga, limpie y cargue los datos "
            "del Excel. Sobre este repositorio se desarrollará un reporte interactivo en Power BI con los "
            "KPIs de finanzas y ventas, y se aplicará la técnica de agrupamiento K-means para segmentar a "
            "los clientes. El Datawarehouse contribuye a la solución porque:")
    viñeta(doc, " los datos limpios y validados se almacenan una sola vez, y todos los reportes "
           "calculan los KPIs con las mismas reglas.", "Centraliza y estandariza:")
    viñeta(doc, " el modelo estrella permite analizar cualquier medida por fecha, cliente, producto, "
           "geografía y modo de envío con consultas simples y rápidas.", "Facilita el análisis multidimensional:")
    viñeta(doc, " el proceso ETL corrige o marca las inconsistencias antes de que lleguen a los "
           "reportes.", "Mejora la calidad:")
    viñeta(doc, " cuando se reciban datos nuevos solo debe ejecutarse el ETL y el tablero se actualiza "
           "automáticamente.", "Es escalable:")

    doc.add_heading("1.5 Objetivos", 2)
    parrafo(doc, "Diseñar e implementar un Datawarehouse que integre los datos de Global Electronic "
            "Retail y un reporte interactivo de KPIs que apoye la toma de decisiones de las áreas de "
            "finanzas y ventas.", "Objetivo general: ")
    parrafo(doc, "", "Objetivos específicos:")
    for o in ("Diseñar un modelo dimensional (esquema estrella) a partir del archivo Excel.",
              "Desarrollar un proceso ETL que limpie, transforme y cargue los datos en PostgreSQL.",
              "Definir y calcular KPIs financieros y comerciales justificados.",
              "Construir un tablero interactivo en Power BI para ambas áreas.",
              "Segmentar a los clientes con K-means para apoyar estrategias comerciales."):
        numerado(doc, o)

    # ---------- 2. Recolección y preparación ----------
    doc.add_heading("2. Recolección y preparación de datos", 1)
    doc.add_heading("2.1 Fuente de datos", 2)
    tabla(doc, ["Característica", "Valor"], [
        ("Archivo", "SAC Retailer Dataset.xlsx (hoja «Retailer Data»)"),
        ("Registros", f"{perfil['filas']:,} líneas de pedido"),
        ("Columnas", "29"),
        ("Periodo", f"{perfil['desde']:%d/%m/%Y} – {perfil['hasta']:%d/%m/%Y}"),
        ("Clientes únicos", f"{perfil['clientes']:,}"),
        ("Productos únicos", f"{perfil['productos']:,}"),
        ("Ciudades únicas", f"{perfil['ciudades']:,} en {perfil['paises']} países"),
        ("Valores nulos", "Ninguno"),
    ], [2.2, 4.3])

    doc.add_heading("2.2 Tipos de datos que se extraerán", 2)
    tabla(doc, ["Grupo", "Columnas del Excel", "Tipo", "Destino en el DW"], [
        ("Identificadores", "Primary ID, Order ID", "Entero", "fact_ventas"),
        ("Fechas", "Order Date, Ship Date", "Fecha", "dim_fecha (rol pedido / envío)"),
        ("Cliente", "Customer ID, Customer Name, Segment", "Texto", "dim_cliente"),
        ("Geografía", "City, State, Country, ISO3, Region, Market", "Texto", "dim_geografia"),
        ("Coordenadas", "Latitude, Longitude", "Decimal", "dim_geografia"),
        ("Producto", "Product ID, Product Name, Sub-Category, Category", "Texto", "dim_producto"),
        ("Envío", "Ship Mode", "Texto", "dim_modo_envio"),
        ("Medidas de venta", "Quantity, Discount, Sales", "Entero / decimal", "fact_ventas"),
        ("Medidas financieras", "Profit, Shipping Costs, Storage Costs, Personnel Cost, "
                                "Selling Costs, Purchasing Costs", "Decimal (moneda)", "fact_ventas"),
    ], [1.3, 2.6, 1.1, 1.5])
    parrafo(doc, "Durante el perfilado se comprobó que la utilidad cumple exactamente la relación "
            "Profit = Sales − (Shipping + Storage + Personnel + Selling + Purchasing Costs). Por ello "
            "el ETL calculará la medida costo_total y validará esta igualdad en cada carga.")

    doc.add_heading("2.3 Problemas de calidad detectados y tratamiento", 2)
    tabla(doc, ["Hallazgo", "Impacto", "Tratamiento en el ETL"], [
        (f"{perfil['envio_antes']} líneas con Ship Date anterior a Order Date",
         "Días de envío negativos", "Se cargan con fecha_inconsistente = TRUE y se excluyen del KPI de "
                                   "tiempo de envío."),
        (f"{perfil['order_id_reusado']:,} Order ID reutilizados con otro cliente o fecha",
         "Conteo de pedidos inflado o reducido", "El grano será Primary ID; un pedido se identifica por "
                                                 "(Order ID, cliente, fecha)."),
        (f"Regiones repetidas en varios mercados ({', '.join(perfil['regiones_ambiguas'])})",
         "Agregaciones erróneas por región", "Jerarquía Mercado → Región → País → Estado → Ciudad."),
        (f"{perfil['descuentos_flotantes']} descuentos con error de punto flotante (p. ej. 0.15000000000000002)",
         "Bandas de descuento duplicadas", "Redondeo a 2 decimales."),
        ("Nombres de producto genéricos («Article 49216»)", "Poca legibilidad",
         "Se conservan; el análisis se hace por categoría y subcategoría."),
        ("Encabezados en inglés con espacios", "Nombres poco prácticos",
         "Renombrado a snake_case en español."),
    ], [2.2, 1.6, 2.7])

    doc.add_heading("2.4 Estructuras del Datawarehouse", 2)
    parrafo(doc, "El Datawarehouse se compone de una tabla de hechos y cinco dimensiones. El grano de la "
            "tabla de hechos es una línea de pedido (un producto dentro de un pedido).")
    tabla(doc, ["Tabla", "Tipo", "Contenido", "Registros aprox."], [
        ("fact_ventas", "Hechos", "Medidas de venta, descuento, costos y utilidad por línea", f"{perfil['filas']:,}"),
        ("dim_fecha", "Dimensión (rol doble)", "Calendario 2019–2021 con jerarquía año–trimestre–mes–día", "1,096"),
        ("dim_cliente", "Dimensión", "Cliente, segmento y clúster K-means", f"{perfil['clientes']:,}"),
        ("dim_producto", "Dimensión", "Producto, subcategoría y categoría", f"{perfil['productos']:,}"),
        ("dim_geografia", "Dimensión", "Ciudad, estado, país, región, mercado y coordenadas", f"{perfil['ciudades']:,}"),
        ("dim_modo_envio", "Dimensión", "Basic, Standard, Express, Prime", "4"),
    ], [1.4, 1.4, 2.6, 1.1])

    # ---------- 3. KPIs ----------
    doc.add_heading("3. Definición de KPIs", 1)
    parrafo(doc, "Los KPIs se seleccionaron a partir de las preguntas que cada área debe responder y de "
            "las columnas disponibles en el archivo. Los valores de referencia se calcularon sobre los "
            "datos originales y servirán para validar el Datawarehouse en la Fase 2.")

    doc.add_heading("3.1 KPIs de finanzas", 2)
    tabla(doc, ["KPI", "Fórmula", "Referencia", "Justificación"], [
        ("Margen de utilidad neta", "Σ utilidad / Σ ventas × 100",
         f"2019: {margen19:.2f} %\n2020: {margen20:.2f} %",
         "Mide cuánto de cada dólar vendido se convierte en ganancia; es el indicador principal de "
         "salud financiera."),
        ("Utilidad neta", "Σ utilidad", f"2019: {dinero(y19['utilidad'])}\n2020: {dinero(y20['utilidad'])}",
         "Resultado absoluto del negocio; permite fijar metas de ganancia por mercado y categoría."),
        ("Rentabilidad sobre costos (ROI operativo)", "Σ utilidad / Σ costo_total × 100",
         f"{(y19['utilidad'] + y20['utilidad']) / (y19['ventas'] + y20['ventas'] - y19['utilidad'] - y20['utilidad']) * 100:.2f} %",
         "Indica el retorno obtenido por cada dólar invertido en operar; sustituye a la rentabilidad de "
         "inversiones, ya que el archivo no contiene montos de inversión."),
        ("Estructura de costos", "Σ costo_rubro / Σ ventas × 100",
         "\n".join(f"{k.split()[0]}: {v:.1f} %" for k, v in perfil["pct_costos"].items()),
         "Identifica el rubro que más consume el ingreso para priorizar acciones de reducción de costos."),
        ("% de líneas con pérdida", "Nº líneas con utilidad < 0 / Nº líneas × 100",
         f"{perfil['perdidas'] / perfil['filas'] * 100:.1f} %",
         "Señala operaciones no rentables y, cruzado con el descuento, revela políticas de precio dañinas."),
    ], [1.3, 1.6, 1.3, 2.3])

    doc.add_heading("3.2 KPIs de ventas", 2)
    tabla(doc, ["KPI", "Fórmula", "Referencia", "Justificación"], [
        ("Ventas totales", "Σ ventas", f"2019: {dinero(y19['ventas'])}\n2020: {dinero(y20['ventas'])}",
         "Volumen de ingresos; base para medir metas comerciales por mercado, categoría y periodo."),
        ("Crecimiento de ventas (interanual)", "(Ventas año actual / Ventas año anterior − 1) × 100",
         f"{crec:.2f} %", "Mide si la empresa se expande; permite detectar mercados o categorías estancados."),
        ("Ventas por cliente", "Σ ventas / Nº clientes distintos",
         f"2019: {dinero(y19['ventas'] / y19['clientes'])}\n2020: {dinero(y20['ventas'] / y20['clientes'])}",
         "Refleja el valor de cada cliente y ayuda a orientar estrategias de fidelización por segmento."),
        ("Ticket promedio por pedido", "Σ ventas / Nº pedidos",
         f"2019: {dinero(y19['ventas'] / y19['pedidos'])}\n2020: {dinero(y20['ventas'] / y20['pedidos'])}",
         "Indica cuánto compra un cliente en cada transacción; sirve para evaluar promociones de venta cruzada."),
        ("Unidades vendidas", "Σ cantidad", f"2019: {y19['unidades']:,}\n2020: {y20['unidades']:,}",
         "Mide la demanda física por producto, útil para planificar inventario y compras."),
        ("Descuento promedio", "Promedio ponderado del descuento", "Por calcular en Fase 2",
         "Controla la política comercial y su efecto sobre el margen."),
    ], [1.3, 1.6, 1.3, 2.3])

    doc.add_heading("3.3 KPIs sugeridos que no aplican", 2)
    parrafo(doc, "La tasa de conversión requiere datos de visitas o cotizaciones, y el flujo de caja "
            "requiere fechas de cobro y pago; ninguno de estos datos existe en el archivo, por lo que se "
            "reemplazaron por la rentabilidad sobre costos y el ticket promedio, que sí pueden calcularse "
            "de forma confiable.")

    # ---------- 4. Herramientas ----------
    doc.add_heading("4. Selección de herramientas y técnicas de minería de datos", 1)
    parrafo(doc, "Se seleccionaron dos herramientas/técnicas de la lista propuesta: Power BI y "
            "agrupamiento con K-means.")
    doc.add_heading("4.1 Power BI", 2)
    viñeta(doc, " se conecta directamente a PostgreSQL y permite construir el tablero interactivo de "
           "KPIs con filtros por año, mercado, categoría, segmento y modo de envío.", "Uso:")
    viñeta(doc, " es gratuito en su versión Desktop, tiene una curva de aprendizaje baja, las medidas "
           "DAX permiten implementar cada KPI una sola vez y reutilizarlo en todos los visuales, e incluye "
           "mapas que aprovechan la latitud y longitud del dataset.", "Justificación:")
    viñeta(doc, " cubre la necesidad de reportes que con otras opciones exigiría combinar Cubos OLAP "
           "y Reporting Services, lo que implica instalar y configurar SQL Server Analysis Services y "
           "Reporting Services.", "Ventaja frente a las alternativas:")
    doc.add_heading("4.2 Agrupamiento con K-means", 2)
    viñeta(doc, " segmentar a los clientes según su comportamiento de compra con las variables "
           "recencia, frecuencia, monto total, margen promedio y descuento promedio (modelo RFM ampliado).",
           "Uso:")
    viñeta(doc, " el área de ventas necesita saber qué clientes son más valiosos y cuáles compran "
           "solo con descuentos de alto costo. K-means es un algoritmo simple, rápido y fácil de "
           "interpretar; el número de grupos se elegirá con el método del codo.", "Justificación:")
    viñeta(doc, " se implementará con Python (scikit-learn) y el clúster resultante se guardará en "
           "la columna cluster_kmeans de dim_cliente, para que Power BI lo use como filtro.",
           "Implementación:")
    doc.add_heading("4.3 Herramientas de apoyo", 2)
    tabla(doc, ["Herramienta", "Uso"], [
        ("PostgreSQL", "Motor del Datawarehouse"),
        ("Python (pandas, openpyxl, SQLAlchemy)", "Proceso ETL"),
        ("Python (scikit-learn)", "Modelo K-means"),
        ("Power BI Desktop", "Modelo semántico, medidas DAX y tablero"),
        ("Git / GitHub", "Control de versiones del código y la documentación"),
    ], [2.6, 3.9])

    # ---------- 5. Diseño ----------
    doc.add_heading("5. Especificación del diseño del Datawarehouse", 1)
    doc.add_heading("5.1 Modelo dimensional (esquema estrella)", 2)
    parrafo(doc, "Se eligió un esquema estrella porque es el modelo más simple de consultar y el que "
            "Power BI recomienda: cada dimensión se une con la tabla de hechos mediante una sola relación "
            "uno a muchos (1:N). La dimensión de fecha cumple un rol doble (fecha de pedido y fecha de "
            "envío).")
    doc.add_picture(str(DIAGRAMS / "star_schema.png"), width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    centrado("Figura 1. Diagrama del esquema estrella.", 9, cursiva=True)

    doc.add_heading("5.2 Diccionario de la tabla de hechos", 2)
    tabla(doc, ["Columna", "Tipo", "Origen", "Descripción"], [
        ("venta_key", "BIGSERIAL (PK)", "Generada", "Clave sustituta"),
        ("primary_id", "INTEGER", "Primary ID", "Trazabilidad con el Excel"),
        ("order_id", "INTEGER", "Order ID", "Dimensión degenerada"),
        ("fecha_pedido_key / fecha_envio_key", "INTEGER (FK)", "Order / Ship Date", "AAAAMMDD"),
        ("cliente_key, producto_key, geografia_key, modo_envio_key", "INTEGER (FK)", "Búsqueda", "Claves de dimensiones"),
        ("cantidad", "SMALLINT", "Quantity", "Unidades vendidas"),
        ("descuento", "NUMERIC(4,2)", "Discount", "Porcentaje (0–1)"),
        ("ventas", "NUMERIC(12,2)", "Sales", "Ingreso de la línea"),
        ("costo_envio … costo_compra", "NUMERIC(12,2)", "Cinco columnas de costo", "Costos por rubro"),
        ("costo_total", "NUMERIC(12,2)", "Calculada", "Suma de los cinco costos"),
        ("utilidad", "NUMERIC(12,2)", "Profit", "Ventas − costo_total"),
        ("dias_envio", "SMALLINT", "Calculada", "Ship Date − Order Date"),
        ("fecha_inconsistente", "BOOLEAN", "Calculada", "TRUE si el envío es anterior al pedido"),
    ], [2.3, 1.3, 1.4, 1.5])

    doc.add_heading("5.3 Jerarquías", 2)
    viñeta(doc, " Año → Trimestre → Mes → Día", "Fecha:")
    viñeta(doc, " Categoría → Subcategoría → Producto", "Producto:")
    viñeta(doc, " Mercado → Región → País → Estado → Ciudad", "Geografía:")
    viñeta(doc, " Segmento → Clúster → Cliente", "Cliente:")

    doc.add_heading("5.4 Proceso ETL (Extract, Transform, Load)", 2)
    parrafo(doc, "El ETL se implementará en Python y se ejecutará con un solo comando. Sus etapas son:")
    parrafo(doc, "", "Extracción")
    for t in ("Leer la hoja «Retailer Data» del archivo Excel con pandas y openpyxl.",
              "Validar que existan las 29 columnas esperadas y registrar el número de filas leídas."):
        viñeta(doc, t)
    parrafo(doc, "", "Transformación")
    for t in ("Renombrar las columnas a snake_case en español y asignar tipos (fecha, entero, decimal).",
              "Limpiar espacios en blanco en textos y redondear Discount a 2 decimales y montos a 2 decimales.",
              "Calcular costo_total, dias_envio y la bandera fecha_inconsistente.",
              "Validar la regla utilidad = ventas − costo_total; las filas que no cumplan se registran en un log.",
              "Generar la tabla de calendario (dim_fecha) para todo el rango de fechas de pedido y envío.",
              "Obtener los valores únicos de cada dimensión (drop_duplicates sobre su clave natural)."):
        viñeta(doc, t)
    parrafo(doc, "", "Carga")
    for t in ("Ejecutar el script DDL (sql/dw_schema.sql) para crear el esquema dw.",
              "Cargar primero las dimensiones y recuperar sus claves sustitutas.",
              "Reemplazar las claves naturales por claves sustitutas en el DataFrame de hechos y cargar "
              "fact_ventas dentro de una transacción.",
              "Verificar la carga: número de filas, suma de ventas y suma de utilidad deben coincidir con el Excel."):
        viñeta(doc, t)
    parrafo(doc, "Flujo resumido: Excel → pandas (extracción) → limpieza y cálculo (transformación) → "
            "dimensiones → tabla de hechos en PostgreSQL (carga) → Power BI y K-means (consumo).")

    # ---------- 6. Cronograma ----------
    doc.add_heading("6. Cronograma de trabajo (fase de implementación)", 1)
    tabla(doc, ["#", "Actividad", "Inicio", "Fin", "Días", "Responsable"], [
        (i + 1, act, dt.date.fromisoformat(a).strftime("%d/%m/%Y"), dt.date.fromisoformat(b).strftime("%d/%m/%Y"),
         (dt.date.fromisoformat(b) - dt.date.fromisoformat(a)).days + 1, resp)
        for i, (act, a, b, resp) in enumerate(CRONOGRAMA)
    ], [0.3, 2.8, 0.9, 0.9, 0.5, 1.1])
    doc.add_picture(str(DIAGRAMS / "cronograma.png"), width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    centrado("Figura 2. Diagrama de Gantt de la fase de implementación.", 9, cursiva=True)

    # ---------- 7. Conclusión ----------
    doc.add_heading("7. Conclusión", 1)
    parrafo(doc, "La propuesta transforma un archivo plano de más de veinte mil filas en un "
            "Datawarehouse con esquema estrella, limpio y validado, que permitirá a las áreas de finanzas "
            "y ventas consultar indicadores confiables desde un único tablero en Power BI. La combinación "
            "de KPIs financieros (margen, rentabilidad y estructura de costos) y comerciales (crecimiento, "
            "ventas por cliente y ticket promedio), junto con la segmentación de clientes mediante K-means, "
            "brinda una base sólida para la toma de decisiones informadas en la siguiente fase.")

    OUTPUT.mkdir(exist_ok=True)
    destino = OUTPUT / "DMD941_Proyecto_Fase_1.docx"
    doc.save(destino)
    return destino


if __name__ == "__main__":
    DIAGRAMS.mkdir(exist_ok=True)
    perfil = perfilar(XLSX)
    diagrama_estrella(DIAGRAMS / "star_schema.png")
    diagrama_gantt(DIAGRAMS / "cronograma.png")
    print("Documento generado:", construir(perfil))
